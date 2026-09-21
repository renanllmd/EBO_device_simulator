import io
import os
import threading
from datetime import datetime

from openpyxl import Workbook, load_workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


# =============================================================================
# CORES
# =============================================================================

AZUL_ESCURO  = colors.HexColor("#1B3A5C")
AZUL_MEDIO   = colors.HexColor("#2E6DA4")
CINZA_HEADER = colors.HexColor("#F2F6FA")
CINZA_LINHA  = colors.HexColor("#F8FAFB")
CINZA_BORDA  = colors.HexColor("#C8D8E8")
BRANCO       = colors.white
PRETO        = colors.HexColor("#1A1A1A")
CINZA_RODAPE = colors.HexColor("#888888")


# =============================================================================
# ESTILOS
# =============================================================================

def _estilos():
    return {

        "cell": ParagraphStyle(
            "cell",
            fontName="Helvetica",
            fontSize=6.5,
            leading=8,
            textColor=PRETO,
            alignment=TA_LEFT,
        ),

        "cc": ParagraphStyle(
            "cc",
            fontName="Helvetica",
            fontSize=6.5,
            leading=8,
            textColor=PRETO,
            alignment=TA_CENTER,
        ),

        "msg": ParagraphStyle(
            "msg",
            fontName="Helvetica",
            fontSize=6,
            leading=7.5,
            textColor=PRETO,
            alignment=TA_LEFT,
        ),

        "ts": ParagraphStyle(
            "ts",
            fontName="Helvetica",
            fontSize=6,
            leading=7.5,
            textColor=PRETO,
            alignment=TA_CENTER,
        ),

        "hdr": ParagraphStyle(
            "hdr",
            fontName="Helvetica-Bold",
            fontSize=6.5,
            leading=8,
            textColor=BRANCO,
            alignment=TA_CENTER,
        ),

        "title": ParagraphStyle(
            "title",
            fontName="Helvetica-Bold",
            fontSize=14,
            textColor=BRANCO,
            alignment=TA_LEFT,
        ),

        "sub": ParagraphStyle(
            "sub",
            fontName="Helvetica",
            fontSize=8,
            textColor=BRANCO,
            alignment=TA_LEFT,
        ),

        "meta": ParagraphStyle(
            "meta",
            fontName="Helvetica",
            fontSize=8,
            textColor=AZUL_ESCURO,
            alignment=TA_LEFT,
        ),

        "metav": ParagraphStyle(
            "metav",
            fontName="Helvetica-Bold",
            fontSize=8,
            textColor=AZUL_MEDIO,
            alignment=TA_LEFT,
        ),
    }

# =============================================================================
# HELPERS PDF
# =============================================================================

_HEADERS    = ["Timestamp", "Dispositivo", "Network", "Ponto", "End.", "FC", "Bit", "IP", "Valor", "Mensagem"]
_COL_WIDTHS = [38*mm, 42*mm, 32*mm, 48*mm, 14*mm, 10*mm, 9*mm, 24*mm, 14*mm, 57*mm]


def _fmt_ts(val):
    try:
        return datetime.strptime(str(val)[:19], "%Y-%m-%d %H:%M:%S").strftime("%d/%m/%Y\n%H:%M:%S")
    except Exception:
        return str(val) if val else ""


def _bloco_cabecalho(nome_arquivo, total, st):
    """Faixa azul escura + linha de metadados."""
    titulo = Table(
        [[Paragraph("EBO Device Simulator", st["title"]),
          Paragraph("Relatório de Auditoria", st["sub"])]],
        colWidths=[120*mm, 130*mm],
    )
    titulo.setStyle(TableStyle([
        ("BACKGROUND",    (0, 0), (-1, -1), AZUL_ESCURO),
        ("LEFTPADDING",   (0, 0), (-1, -1), 8),
        ("RIGHTPADDING",  (0, 0), (-1, -1), 8),
        ("TOPPADDING",    (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
    ]))

    agora = datetime.now().strftime("%d/%m/%Y %H:%M")
    meta = Table(
        [[
            Paragraph("Data de geração:", st["meta"]),
            Paragraph(agora, st["metav"]),
            Paragraph("   "),
            Paragraph("Total de eventos:", st["meta"]),
            Paragraph(str(total), st["metav"]),
            Paragraph("   "),
            Paragraph("Arquivo:", st["meta"]),
            Paragraph(nome_arquivo, st["metav"]),
        ]],
        colWidths=[28*mm, 38*mm, 6*mm, 30*mm, 16*mm, 6*mm, 20*mm, 106*mm],
    )
    meta.setStyle(TableStyle([
        ("BACKGROUND",    (0, 0), (-1, -1), CINZA_HEADER),
        ("LEFTPADDING",   (0, 0), (-1, -1), 5),
        ("RIGHTPADDING",  (0, 0), (-1, -1), 2),
        ("TOPPADDING",    (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW",     (0, 0), (-1, -1), 0.5, AZUL_MEDIO),
    ]))

    return [titulo, meta, Spacer(1, 4)]


def _bloco_tabela(rows_raw, st):
    """Tabela de dados com zebra striping e larguras fixas."""
    header_row = [Paragraph(h, st["hdr"]) for h in _HEADERS]
    table_data = [header_row]

    for raw in rows_raw:
        table_data.append([
            Paragraph(_fmt_ts(raw[0]),        st["ts"]),
            Paragraph(str(raw[1] or ""),      st["cell"]),
            Paragraph(str(raw[2] or ""),      st["cell"]),
            Paragraph(str(raw[3] or ""),      st["cell"]),
            Paragraph(str(raw[4] or ""),      st["cc"]),
            Paragraph(str(raw[5] or ""),      st["cc"]),
            Paragraph(str(raw[6] or ""),      st["cc"]),
            Paragraph(str(raw[7] or ""),      st["cc"]),
            Paragraph(str(raw[8] or ""),      st["cc"]),
            Paragraph(str(raw[9] or ""),      st["msg"]),
        ])

    style_cmds = [
        ("BACKGROUND",    (0, 0), (-1, 0),  AZUL_MEDIO),
        ("LINEBELOW",     (0, 0), (-1, 0),  1.2, AZUL_ESCURO),
        ("BOX",           (0, 0), (-1, -1), 0.7, AZUL_MEDIO),
        ("INNERGRID",     (0, 0), (-1, -1), 0.3, CINZA_BORDA),
        ("LEFTPADDING",   (0, 0), (-1, -1), 3),
        ("RIGHTPADDING",  (0, 0), (-1, -1), 3),
        ("TOPPADDING",    (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("VALIGN",        (0, 0), (-1, -1), "TOP"),
    ]
    for i in range(1, len(table_data)):
        bg = CINZA_LINHA if i % 2 == 0 else BRANCO
        style_cmds.append(("BACKGROUND", (0, i), (-1, i), bg))

    tbl = Table(table_data, colWidths=_COL_WIDTHS, repeatRows=1)
    tbl.setStyle(TableStyle(style_cmds))
    return tbl


def _rodape(canvas, doc):
    canvas.saveState()
    w, _h = landscape(A4)
    canvas.setStrokeColor(AZUL_MEDIO)
    canvas.setLineWidth(0.5)
    canvas.line(15*mm, 12*mm, w - 15*mm, 12*mm)
    canvas.setFont("Helvetica", 6)
    canvas.setFillColor(CINZA_RODAPE)
    canvas.drawString(15*mm, 8.5*mm, "EBO Device Simulator — Auditoria Modbus")
    canvas.drawRightString(w - 15*mm, 8.5*mm, f"Página {doc.page}")
    canvas.restoreState()


# =============================================================================
# EXCEL LOGGER — Auditoria de escritas
# =============================================================================

class ExcelLogger:
    def __init__(self, base_dir: str):
        self.file_path = os.path.join(
            base_dir,
            f"EBO_device_simulator_log_{datetime.now().strftime('%d_%m_%Y')}.xlsx"
        )
        self._lock = threading.Lock()
        self._criar_arquivo()

    # ── criação do arquivo ────────────────────────────────────────────────────

    def _criar_arquivo(self):
        if os.path.exists(self.file_path):
            return
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Writes"
        ws.append([
            "Timestamp",
            "Dispositivo",
            "Network",
            "Ponto",
            "Endereco",
            "FunctionCode",
            "Bit",
            "IP",
            "Valor",
            "Mensagem de envio",
        ])
        wb.save(self.file_path)

    # ── escrita de linha ──────────────────────────────────────────────────────

    def adicionar_linha(self, root, net, ponto, endereco, fc, bit, ip, valor, mensagem):
        with self._lock:
            wb = load_workbook(self.file_path)
            ws = wb["Writes"]
            ws.append([
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                root,
                net,
                ponto,
                endereco,
                fc,
                bit if bit is not None else "",
                ip if ip else "LOCAL",
                valor,
                mensagem,
            ])
            wb.save(self.file_path)

    # ── geração do PDF ────────────────────────────────────────────────────────

    def gerar_pdf(self):
        """Lê o XLSX atual e retorna (buffer_pdf, nome_arquivo) sem escrever em disco."""
        with self._lock:
            wb = load_workbook(self.file_path, data_only=True)

        ws = wb["Writes"]
        rows_raw = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
        total = len(rows_raw)
        nome_arquivo = os.path.basename(self.file_path).replace(".xlsx", "")

        st = _estilos()
        story = _bloco_cabecalho(nome_arquivo, total, st)
        story.append(_bloco_tabela(rows_raw, st))

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            leftMargin=15*mm, rightMargin=15*mm,
            topMargin=12*mm,  bottomMargin=18*mm,
        )
        doc.build(story, onFirstPage=_rodape, onLaterPages=_rodape)
        buffer.seek(0)
        return buffer, nome_arquivo
